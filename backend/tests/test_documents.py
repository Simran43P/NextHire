"""
Getting text out of an upload.

Format is decided by the bytes, never by the filename. An extension is
attacker-controlled input, and a `.pdf` that is really a zip is the oldest
trick there is.
"""

import io

import docx
import pytest

import documents


@pytest.fixture
def docx_bytes():
    def build(paragraphs=None, table_rows=None) -> bytes:
        document = docx.Document()
        for text in paragraphs or ["Asha Menon", "Skills: Python, React, SQL"]:
            document.add_paragraph(text)
        if table_rows:
            table = document.add_table(rows=len(table_rows), cols=2)
            for index, (left, right) in enumerate(table_rows):
                table.rows[index].cells[0].text = left
                table.rows[index].cells[1].text = right
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    return build


class TestDetect:
    def test_a_pdf_is_recognised(self):
        assert documents.detect(b"%PDF-1.7\n...") == "pdf"

    def test_a_docx_is_recognised(self, docx_bytes):
        assert documents.detect(docx_bytes()) == "docx"

    def test_plain_text_is_not(self):
        assert documents.detect(b"Asha Menon\nPython developer") is None

    def test_an_empty_file_is_not(self):
        assert documents.detect(b"") is None

    def test_the_filename_is_never_consulted(self, docx_bytes):
        # The function does not take one, which is the point.
        assert documents.detect(docx_bytes()) == "docx"


class TestExtract:
    def test_a_pdf_round_trips(self, pdf_bytes):
        text, kind = documents.extract_text(pdf_bytes())
        assert kind == "pdf"
        assert "Asha Menon" in text

    def test_a_docx_round_trips(self, docx_bytes):
        text, kind = documents.extract_text(docx_bytes())
        assert kind == "docx"
        assert "Asha Menon" in text
        assert "Python" in text

    def test_docx_tables_are_read(self, docx_bytes):
        """
        Plenty of resumes lay themselves out in tables, and their text lives
        nowhere else in the document body.
        """
        content = docx_bytes(
            paragraphs=["Asha Menon"],
            table_rows=[("Zeta Labs", "Software Intern, 2024")],
        )
        text, _ = documents.extract_text(content)
        assert "Zeta Labs" in text
        assert "Software Intern, 2024" in text

    def test_plain_text_is_refused(self):
        with pytest.raises(documents.UnsupportedDocument):
            documents.extract_text(b"Asha Menon, Python developer")

    def test_a_zip_that_is_not_a_word_file_is_refused(self):
        import zipfile

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("hello.txt", "not a document")

        # A .pages file, an .odt, a renamed archive - all unsupported rather
        # than corrupt, and none of them should reach the model.
        with pytest.raises(documents.UnsupportedDocument):
            documents.extract_text(buffer.getvalue())

    def test_a_damaged_pdf_is_reported_as_unreadable(self):
        with pytest.raises(documents.UnreadableDocument):
            documents.extract_text(b"%PDF-1.4 truncated garbage")


class TestUploadRoute:
    def test_a_docx_resume_is_accepted(self, client, mock_model, docx_bytes, sample_profile):
        mock_model(extract={"profile": sample_profile, "gaps": []})
        response = client.post(
            "/api/parse-resume",
            files={
                "file": (
                    "resume.docx",
                    docx_bytes(
                        paragraphs=[
                            "Asha Menon",
                            "asha.menon@example.com | Kochi",
                            "Skills: Python, React, SQL, FastAPI, Git",
                            "Education: B.Tech Computer Science, CUSAT, 2026",
                            "Projects: Resume Ranker, built with Python and scikit-learn "
                            "and deployed using Streamlit for a university project.",
                        ]
                    ),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert response.status_code == 200, response.text
        assert response.json()["profile"]["name"] == "Asha Menon"

    def test_a_renamed_text_file_is_refused(self, client, mock_model):
        mock_model()
        response = client.post(
            "/api/parse-resume",
            files={"file": ("resume.pdf", b"just some text", "application/pdf")},
        )
        assert response.status_code == 400
        assert response.json()["detail"]["code"] == "unsupported_format"

    def test_an_empty_docx_is_refused(self, client, mock_model, docx_bytes):
        mock_model()
        response = client.post(
            "/api/parse-resume",
            files={"file": ("resume.docx", docx_bytes(paragraphs=[""]), "application/octet-stream")},
        )
        assert response.status_code == 400
        assert response.json()["detail"]["code"] == "no_text_in_document"
